"""Behavioral regressions for boxed log ingestion and exported workbook semantics."""
import importlib.util
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "scripts/aisbench_to_excel.py"
spec = importlib.util.spec_from_file_location("converter", SCRIPT)
converter = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = converter
spec.loader.exec_module(converter)
OFFICIAL = (ROOT / "tests/official_excerpt.log").read_text(encoding="utf-8")


def perf(rows, header="| Performance Parameters | Stage | Average | P95 | P99 | N |"):
    return header + "\n" + "\n".join("| " + " | ".join(row) + " |" for row in rows) + "\n"


def common(maximum="32", stage="total"):
    return f"| Common Metric | Stage | Value |\n| Max Concurrency | {stage} | {maximum} |\n| Concurrency | {stage} | 29.5 |\n| Output Token Throughput | {stage} | 8,000 token/s |\n"


def report(stage="total"):
    return perf([["TTFT", stage, "1.5 s", "1,600 ms", "1700000 us", "100"],
                 ["TPOT", stage, "30 ms", "35 ms", "4e1 ms", "100"],
                 ["InputTokens", stage, "131070", "131072", "131072", "100"]]) + common(stage=stage)


class ParseTests(unittest.TestCase):
    def rows(self, text, source="run.log", defaults=None):
        runs, issues = converter.parse_text(text, source, defaults)
        columns, rows, details = converter.build_rows(runs, issues)
        return runs, issues, columns, rows, details

    def test_official_excerpt_expected_values(self):
        runs, issues, columns, rows, details = self.rows(OFFICIAL, "c1_in2k_out512.log")
        self.assertEqual(len(runs), 1)
        self.assertEqual(len(rows), 1)
        self.assertEqual(issues, [])
        row = rows[0]
        self.assertEqual(row["模型"], "vllm-api-stream-chat")
        self.assertEqual(row["数据集"], "demo_gsm8k")
        self.assertEqual(row["请求/TTFT/Average (ms)"], 2006)
        self.assertEqual(row["配置输入长度(token)"], 2048)
        self.assertEqual(row["实测平均输入(token)"], 1512.5)
        self.assertEqual(row["输出吞吐(token/s)"], 41.622)
        self.assertEqual(row["请求/OutputTokenThroughput/Average (token/s)"], 41.6254)
        self.assertEqual(row["整体/Total Output Tokens/Value (token)"], 4096)

    def test_multiple_experiments_in_one_log(self):
        text = '# AISBENCH_META {"concurrency":32,"input_len":"128k"}\n' + report()
        text += '# AISBENCH_META {"concurrency":64,"input_len":"64k"}\n' + report().replace("| 32 |", "| 64 |")
        runs, issues, _, rows, _ = self.rows(text)
        self.assertEqual(len(runs), 2)
        self.assertEqual([r["配置输入长度(token)"] for r in rows], [131072, 65536])
        self.assertEqual([r["配置并发"] for r in rows], [32, 64])
        self.assertEqual(issues, [])

    def test_marker_does_not_leak_to_next_run(self):
        text = '# AISBENCH_META {"input_len":1024}\n' + OFFICIAL + OFFICIAL
        _, _, _, rows, _ = self.rows(text)
        self.assertEqual([r["配置输入长度(token)"] for r in rows], [1024, None])

    def test_total_and_stable_not_mixed(self):
        text = perf([["TTFT", "total", "1 s", "1100 ms", "1.2 s", "100"],
                     ["TTFT", "stable", ".1 s", "110 ms", ".12 s", "80"]]) + common()
        runs, _, _, rows, _ = self.rows(text)
        self.assertEqual(len(runs), 1)
        self.assertEqual([r["阶段"] for r in rows], ["total", "stable"])
        self.assertEqual([r["请求/TTFT/Average (ms)"] for r in rows], [1000, 100])
        self.assertEqual([r["配置并发"] for r in rows], [32, 32])
        self.assertIsNone(rows[1]["实际平均并发"])

    def test_time_units_and_custom_percentile(self):
        _, _, _, rows, _ = self.rows(report(), "c32_in128k.log")
        self.assertEqual(rows[0]["请求/TTFT/Average (ms)"], 1500)
        self.assertEqual(rows[0]["请求/TTFT/P95 (ms)"], 1600)
        self.assertEqual(rows[0]["请求/TTFT/P99 (ms)"], 1700)
        self.assertEqual(rows[0]["请求/TPOT/P99 (ms)"], 40)

    def test_ansi_cr_and_log_prefix(self):
        decorated = "\r".join("[INFO] \x1b[32m" + line + "\x1b[0m" for line in report().splitlines())
        runs, _, _, rows, _ = self.rows(decorated)
        self.assertEqual(len(runs), 1)
        self.assertEqual(rows[0]["请求/TTFT/Average (ms)"], 1500)

    def test_legacy_two_column_and_header_units(self):
        text = "| Performance Parameters | Average (s) | N |\n| TTFT | 2 | 4 |\n| Common Metric | Value |\n| Max Concurrency | 4 |\n"
        _, _, _, rows, _ = self.rows(text)
        self.assertEqual(rows[0]["请求/TTFT/Average (ms)"], 2000)
        self.assertEqual(rows[0]["请求/TTFT/N"], 4)
        self.assertEqual(rows[0]["阶段"], "total")

    def test_reversed_table_order_and_multiple_blocks(self):
        runs, _, _, rows, _ = self.rows(common() + perf([["TTFT", "total", "20 ms", "22 ms", "24 ms", "10"]]))
        self.assertEqual(len(runs), 1)
        self.assertEqual(len(rows), 1)
        runs, _, _, rows, _ = self.rows(report() + report())
        self.assertEqual(len(runs), 2)
        self.assertEqual(len(rows), 2)

    def test_missing_values_keep_raw(self):
        text = perf([["TTFT", "total", "N/A", "nan", "-", "0"]])
        _, issues, _, rows, details = self.rows(text)
        self.assertIsNone(rows[0]["请求/TTFT/Average"])
        self.assertEqual(details[0][7], "N/A")
        self.assertTrue(any("缺失值" in warning[2] for warning in issues))

    def test_unknown_latency_unit_not_in_ms_column(self):
        _, issues, _, rows, _ = self.rows(perf([["TTFT", "total", "2", "3", "4", "10"]]))
        self.assertEqual(rows[0]["请求/TTFT/Average"], 2)
        self.assertNotIn("请求/TTFT/Average (ms)", rows[0])
        self.assertTrue(any("未给单位" in item[2] for item in issues))

    def test_duplicate_rows_keep_both_details(self):
        text = perf([["TTFT", "total", "1 ms", "2 ms", "3 ms", "10"],
                     ["TTFT", "total", "9 ms", "10 ms", "11 ms", "10"]])
        _, issues, _, rows, details = self.rows(text)
        self.assertEqual(rows[0]["请求/TTFT/Average (ms)"], 1)
        self.assertEqual(len(details), 8)
        self.assertTrue(any("重复指标" in i[2] for i in issues))

    def test_truncated_row_warns_without_fabrication(self):
        text = perf([["TTFT", "total", "1 ms", "2 ms", "3 ms", "10"]]) + "| TPOT | total | 4 ms |\n"
        _, issues, _, rows, _ = self.rows(text)
        self.assertFalse(any("TPOT" in k for k in rows[0]))
        self.assertTrue(any("列数不匹配" in i[2] for i in issues))

    def test_actual_vs_configured_concurrency(self):
        _, issues, _, rows, _ = self.rows(report(), "c64_in128k.log")
        self.assertEqual(rows[0]["配置并发"], 32)
        self.assertEqual(rows[0]["实际平均并发"], 29.5)
        self.assertTrue(any("并发冲突" in i[2] for i in issues))

    def test_no_guess_from_measured_length(self):
        _, _, _, rows, _ = self.rows(report())
        self.assertIsNone(rows[0]["配置输入长度(token)"])
        self.assertEqual(rows[0]["实测平均输入(token)"], 131070)

    def test_no_progress_table_as_results(self):
        runs, issues = converter.parse_text("| Task Name | Progress |\n| model/dataset | 100% |\n")
        self.assertEqual(runs, [])
        self.assertTrue(issues)

    def test_metadata_priority_and_invalid_marker(self):
        marker = '# AISBENCH_META {"input_len":"64k"}\n'
        _, _, _, rows, _ = self.rows(marker + report(), "c32_in128k.log", {"input_len": "32k"})
        self.assertEqual(rows[0]["配置输入长度(token)"], 65536)
        _, issues, _, _, _ = self.rows('# AISBENCH_META {"concurency":32}\n' + report())
        self.assertTrue(any("元数据无效" in i[2] for i in issues))

    def test_literal_text_not_formulas_and_numbers_are_numeric(self):
        from openpyxl import load_workbook
        _, issues, columns, rows, details = self.rows(report(), defaults={"model": "=1+1", "input_len": 128})
        with tempfile.TemporaryDirectory() as temp:
            output = Path(temp) / "result.xlsx"
            converter.write_excel(output, columns, rows, details, issues)
            wb = load_workbook(output)
            ws = wb["汇总"]
            self.assertEqual(ws.cell(2, 3).value, "=1+1")
            self.assertEqual(ws.cell(2, 3).data_type, "s")
            self.assertEqual(ws.cell(2, columns.index("配置并发") + 1).data_type, "n")
            self.assertEqual(ws.freeze_panes, "F2")
            self.assertEqual(len(ws.tables), 1)
            self.assertEqual(wb["指标明细"].max_row, len(details) + 1)
            wb.close()

    def test_cli_files_pipe_strict_and_overwrite(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            log = folder / "c32_in128k.log"
            log.write_text(report(), encoding="utf-8")
            out = folder / "result.xlsx"
            def cli(*args, text=None):
                return subprocess.run([sys.executable, str(SCRIPT), *map(str, args)], input=text,
                                      capture_output=True, encoding="utf-8", errors="replace")
            self.assertEqual(cli(folder, "-o", out, "--strict").returncode, 0)
            self.assertEqual(cli(log, "-o", out).returncode, 2)
            self.assertEqual(cli(log, "-o", out, "--force").returncode, 0)
            piped = folder / "pipe.xlsx"
            self.assertEqual(cli("-", "-o", piped, "--input-len", "128k", text=report()).returncode, 0)
            strict = folder / "strict.xlsx"
            self.assertEqual(cli("-", "-o", strict, "--strict", text=report()).returncode, 2)
            self.assertFalse(strict.exists())

    def test_multiple_files_glob_and_recursive(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            (folder / "sub").mkdir()
            (folder / "a.log").write_text(report(), encoding="utf-8")
            (folder / "sub/b.out").write_text(report(), encoding="utf-8")
            (folder / "ignore.xlsx").write_text("skip", encoding="utf-8")
            self.assertEqual(len(converter.collect_inputs([str(folder)], False)), 1)
            self.assertEqual(len(converter.collect_inputs([str(folder)], True)), 2)
            self.assertEqual(len(converter.collect_inputs([str(folder / "*.log"), str(folder / "a.log")], False)), 1)

    def test_encoding_utf16_and_gb18030(self):
        with tempfile.TemporaryDirectory() as temp:
            for encoding in ("utf-8-sig", "utf-16", "gb18030"):
                path = Path(temp) / "log.txt"
                path.write_bytes(("中文\n" + OFFICIAL).encode(encoding))
                self.assertEqual(converter.read_text(path, "auto"), "中文\n" + OFFICIAL)

    def test_numeric_formats(self):
        self.assertEqual(converter.parse_value("1,234.50 tokens/s")[:2], (1234.5, "token/s"))
        self.assertEqual(converter.parse_value("20 %")[:2], (.2, "%"))
        self.assertEqual(converter.parse_value("2 µs")[:2], (.002, "ms"))
        self.assertEqual(converter.parse_value("1e-3 s")[:2], (1, "ms"))
        self.assertEqual(converter.token_count("128k"), 131072)


if __name__ == "__main__":
    unittest.main()
