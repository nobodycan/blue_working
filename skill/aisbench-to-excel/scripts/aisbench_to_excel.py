#!/usr/bin/env python3
"""Convert AISBench boxed terminal reports to XLSX. Python >= 3.9, openpyxl >= 3.1."""
from __future__ import annotations

import argparse
import glob
import json
import math
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

ANSI = re.compile(r"\x1b\][^\x07\x1b]*(?:\x07|\x1b\\)|\x1b\[[0-?]*[ -/]*[@-~]")
NUMBER = r"[+-]?(?:(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?"
TIME_METRICS = {"e2el", "ttft", "tpot", "itl", "tokenizer", "detokenizer", "benchmarkduration"}
TASK = re.compile(r"Performance Results? of task\s*:\s*(.*?)\s*:\s*$", re.I)
META = re.compile(r"(?:#\s*)?AISBENCH_META\s+(\{.*\})\s*$")
MISSING = {"", "-", "--", "n/a", "na", "none", "null", "nan", "inf", "+inf", "-inf", "∞", "—"}
ALIASES = {"totalgeneratedtokens": "Total Output Tokens", "successfulrequests": "Success Requests"}
META_FIELDS = {"concurrency", "input_len", "output_len", "model", "dataset", "tag"}
SUMMARY_BASE = ["结果ID", "阶段", "模型", "数据集", "标签", "配置并发", "实际平均并发",
                "配置输入长度(token)", "配置输出长度(token)", "实测平均输入(token)",
                "实测平均输出(token)", "总请求数", "成功请求数", "失败请求数", "时长(ms)",
                "请求吞吐(req/s)", "输出吞吐(token/s)", "总吞吐(token/s)"]


def compact(text):
    return re.sub(r"[^a-z0-9]", "", text.lower())


def label_unit(label):
    """Only interpret a trailing recognized unit, leaving custom metric labels intact."""
    m = re.search(r"\s*[\[(](ms|us|µs|μs|s|seconds?|milliseconds?|tokens?/s|req/s|requests?/s|%)[\])]\s*$", label, re.I)
    return (label[:m.start()].strip(), m.group(1)) if m else (label.strip(), "")


def parse_value(raw, metric="", unit_hint=""):
    """Return typed value, canonical unit, issue. Raw text is retained separately."""
    text = raw.strip()
    if text.lower() in MISSING:
        return None, unit_hint, "缺失值或非有限值"
    match = re.fullmatch(rf"({NUMBER})\s*(.*?)", text)
    if not match:
        return text, unit_hint, "无法识别为数值，已保留文本"
    number = float(match.group(1).replace(",", ""))
    if not math.isfinite(number):
        return None, unit_hint, "缺失值或非有限值"
    unit = (match.group(2) or unit_hint).strip().lower()
    unit = {"μs": "us", "µs": "us", "sec": "s", "second": "s", "seconds": "s",
            "millisecond": "ms", "milliseconds": "ms", "tokens/s": "token/s",
            "request/s": "req/s", "requests/s": "req/s", "tokens": "token"}.get(unit, unit)
    if unit in {"s", "ms", "us", "ns"}:
        number *= {"s": 1000, "ms": 1, "us": .001, "ns": .000001}[unit]
        unit = "ms"
    elif unit == "%":
        number /= 100
    elif not unit and compact(metric) in TIME_METRICS:
        # Do not silently assume seconds or milliseconds for a unitless latency.
        return number, "", "时延未给单位，未做换算"
    elif not unit and compact(metric) in {"inputtokens", "outputtokens", "totalinputtokens", "totaloutputtokens"}:
        unit = "token"
    elif unit and unit not in {"token", "token/s", "req/s"}:
        return number, unit, "未知单位，未做换算"
    if number.is_integer() and abs(number) < 10**15:
        number = int(number)
    return number, unit, ""


def token_count(value):
    match = re.fullmatch(r"(\d+(?:\.\d+)?)\s*([kKmM]?)", str(value).strip())
    if not match:
        raise ValueError(f"无效长度/并发: {value!r}")
    count = float(match.group(1)) * {"": 1, "k": 1024, "m": 1024**2}[match.group(2).lower()]
    if not count.is_integer() or count < 0:
        raise ValueError(f"必须是非负整数: {value!r}")
    return int(count)


def normalize_meta(meta):
    unknown = set(meta) - META_FIELDS
    if unknown:
        raise ValueError("不支持的元数据字段: " + ", ".join(sorted(unknown)))
    result = dict(meta)
    for key in ("concurrency", "input_len", "output_len"):
        if key in result:
            result[key] = token_count(result[key])
    if result.get("concurrency") == 0:
        raise ValueError("concurrency 必须大于零")
    for key in ("model", "dataset", "tag"):
        if key in result and not isinstance(result[key], str):
            raise ValueError(f"{key} 必须是字符串")
    return result


def filename_meta(source):
    """Read only explicit filename conventions, never arbitrary positional numbers."""
    name = Path(source).stem.lower()
    aliases = {"concurrency": r"(?:concurrency|conc|batch_size|c)",
               "input_len": r"(?:input_len|input|in|i)",
               "output_len": r"(?:output_len|output|out|o)"}
    result = {}
    for key, alias in aliases.items():
        hits = re.findall(rf"(?:^|[_\- ]){alias}[=_\-]?(\d+(?:\.\d+)?[km]?)(?=$|[_\- .])", name)
        if len(set(hits)) == 1:
            result[key] = token_count(hits[0])
    return result


@dataclass
class Metric:
    table: str
    metric: str
    stage: str
    stat: str
    value: object
    unit: str
    raw: str
    line: int


@dataclass
class Run:
    source: str
    index: int
    start: int
    end: int = 0
    task: str = ""
    meta: dict = field(default_factory=dict)
    origins: dict = field(default_factory=dict)
    metrics: list = field(default_factory=list)
    tables: set = field(default_factory=set)


def table_cells(line):
    # AISBench uses tabulate's fancy_grid, grid or pipe; log prefixes may precede rows.
    candidates = [p for p in (line.find("│"), line.find("|"), line.find("┃")) if p >= 0]
    if not candidates:
        return None
    body = line[min(candidates):].strip()
    if body[-1:] not in {"│", "|", "┃"}:
        return None
    return [x.strip() for x in re.split(r"[│|┃]", body)[1:-1]]


def parse_text(text, source="<stdin>", defaults=None):
    """Parse ordered complete reports; unrelated progress/config tables are ignored."""
    defaults = normalize_meta(defaults or {})
    base = filename_meta(source) if source != "<stdin>" else {}
    origins = {k: "文件名" for k in base}
    base.update(defaults)
    origins.update({k: "命令行" for k in defaults})
    runs, issues = [], []
    run = None
    header = None
    kind = ""

    def warn(line, message):
        issues.append((source, line, message))

    def new_run(line, extra=None):
        meta = dict(base)
        meta.update(extra or {})
        provenance = dict(origins)
        provenance.update({k: "AISBENCH_META" for k in (extra or {})})
        return Run(source, len(runs) + 1, line, meta=meta, origins=provenance)

    def finish(line):
        nonlocal run, header, kind
        if run and run.metrics:
            run.end = line
            if run.tables != {"performance", "common"}:
                warn(run.start, "结果只包含一类表格，可能是截断日志；已保留可解析指标")
            runs.append(run)
        run, header, kind = None, None, ""

    cleaned = ANSI.sub("", text).replace("\r\n", "\n").replace("\r", "\n")
    for line_no, line in enumerate(cleaned.splitlines(), 1):
        line = re.sub(r"[^\x08]\x08", "", line)
        marker = META.search(line)
        if marker:
            finish(line_no - 1)
            try:
                extra = json.loads(marker.group(1))
                if not isinstance(extra, dict):
                    raise ValueError("元数据必须为 JSON 对象")
                run = new_run(line_no, normalize_meta(extra))
            except (ValueError, TypeError) as exc:
                warn(line_no, f"元数据无效: {exc}")
            continue
        task = TASK.search(line)
        if task:
            if run and run.metrics:
                finish(line_no - 1)
            if run is None:
                run = new_run(line_no)
            run.task = task.group(1)
            if "/" in run.task:
                model, dataset = run.task.split("/", 1)
                for key, value in (("model", model), ("dataset", dataset)):
                    if key not in run.meta:
                        run.meta[key] = value
                        run.origins[key] = "结果任务标题"
            header, kind = None, ""
            continue
        if re.search(r"Performance Result files? located", line, re.I):
            finish(line_no)
            continue
        cells = table_cells(line)
        if not cells or not any(cells):
            continue
        names = [compact(label_unit(c)[0]) for c in cells]
        new_kind = "performance" if names[0] == "performanceparameters" else (
            "common" if names[0] in {"commonmetric", "commonmetrics"} and "value" in names else "")
        if new_kind:
            if run and new_kind in run.tables:
                finish(line_no - 1)
            if run is None:
                run = new_run(line_no)
            kind, header = new_kind, cells
            run.tables.add(kind)
            continue
        if not header or run is None:
            continue
        if all(not c or re.fullmatch(r"[:=\-+ ]+", c) for c in cells):
            continue
        if len(cells) != len(header):
            warn(line_no, f"表格列数不匹配: 期望 {len(header)}，实际 {len(cells)}；已跳过")
            continue
        names = [compact(label_unit(c)[0]) for c in header]
        metric, metric_unit = label_unit(cells[0])
        metric = ALIASES.get(compact(metric), metric)
        if not metric:
            warn(line_no, "指标名为空，已跳过该行（不猜测被折行的内容）")
            continue
        stage = cells[names.index("stage")] if "stage" in names else "total"
        stage = stage.strip().lower() or "total"
        for column in range(1, len(header)):
            if names[column] == "stage":
                continue
            stat, stat_unit = label_unit(header[column])
            stat = {"average": "Average", "mean": "Average", "avg": "Average",
                    "median": "Median", "p50": "Median", "n": "N", "value": "Value",
                    "min": "Min", "max": "Max"}.get(compact(stat), stat.upper())
            hint = "" if stat == "N" else (stat_unit or metric_unit)
            value, unit, problem = parse_value(cells[column], "" if stat == "N" else metric, hint)
            run.metrics.append(Metric(kind, metric, stage, stat, value, unit, cells[column], line_no))
            if problem:
                warn(line_no, f"{metric}/{stage}/{stat}: {problem} ({cells[column]})")
    finish(len(cleaned.splitlines()))
    if not runs:
        warn(0, "未找到可解析的 AISBench 性能结果表")
    return runs, issues


def metric_column(metric):
    suffix = f" ({metric.unit})" if metric.unit else ""
    group = "整体" if metric.table == "common" else "请求"
    return f"{group}/{metric.metric}/{metric.stat}{suffix}"


def build_rows(runs, issues):
    """One summary row per result/stage; no averaging or deduplication of repeated runs."""
    summary, details = [], []
    all_columns = []
    for seq, run in enumerate(runs, 1):
        result_id = f"R{seq:05d}"
        for stage in dict.fromkeys(m.stage for m in run.metrics):
            metrics = [m for m in run.metrics if m.stage == stage]
            row = dict(zip(SUMMARY_BASE, [None] * len(SUMMARY_BASE)))
            row.update({"结果ID": result_id, "阶段": stage, "模型": run.meta.get("model"),
                        "数据集": run.meta.get("dataset"), "标签": run.meta.get("tag"),
                        "配置并发": run.meta.get("concurrency"),
                        "配置输入长度(token)": run.meta.get("input_len"),
                        "配置输出长度(token)": run.meta.get("output_len")})
            seen = set()
            for m in metrics:
                column = metric_column(m)
                duplicate = column in seen
                details.append([result_id, stage, m.table, m.metric, m.stat, m.value, m.unit,
                                m.raw, run.source, m.line, "重复，汇总采用首个值" if duplicate else ""])
                if duplicate:
                    issues.append((run.source, m.line, f"重复指标 {stage}/{column}；明细全部保留，汇总采用首个值"))
                    continue
                seen.add(column)
                row[column] = m.value
                if column not in all_columns:
                    all_columns.append(column)

            def get(table, metric, stat="Value", unit=None):
                return next((m.value for m in metrics if m.table == table and compact(m.metric) == compact(metric)
                             and m.stat == stat and (unit is None or m.unit == unit)), None)

            actual_max = get("common", "Max Concurrency")
            if actual_max is None:
                run_maxima = {m.value for m in run.metrics if m.table == "common"
                              and compact(m.metric) == "maxconcurrency"
                              and isinstance(m.value, (int, float))}
                if len(run_maxima) == 1:
                    actual_max = next(iter(run_maxima))
            provenance = dict(run.origins)
            if isinstance(actual_max, (int, float)):
                if row["配置并发"] is not None and row["配置并发"] != actual_max:
                    issues.append((run.source, run.start, f"并发冲突: 元数据={row['配置并发']}，Max Concurrency={actual_max}；采用结果表"))
                row["配置并发"] = actual_max
                provenance["concurrency"] = "Common Metric/Max Concurrency"
            for label, metric, unit in [
                ("实际平均并发", "Concurrency", None), ("总请求数", "Total Requests", None),
                ("成功请求数", "Success Requests", None), ("失败请求数", "Failed Requests", None),
                ("时长(ms)", "Benchmark Duration", "ms"), ("请求吞吐(req/s)", "Request Throughput", "req/s"),
                ("输出吞吐(token/s)", "Output Token Throughput", "token/s"),
                ("总吞吐(token/s)", "Total Token Throughput", "token/s")]:
                row[label] = get("common", metric, unit=unit)
            row["实测平均输入(token)"] = get("performance", "InputTokens", "Average")
            row["实测平均输出(token)"] = get("performance", "OutputTokens", "Average")
            for key, label in (("配置并发", "配置并发"), ("配置输入长度(token)", "配置输入长度")):
                if row[key] is None:
                    issues.append((run.source, run.start, f"{stage}: 未知{label}，留空；可用文件名、AISBENCH_META 或命令行补充"))
            row.update({"任务": run.task, "来源文件": run.source, "文件内结果序号": run.index,
                        "起始行": run.start, "结束行": run.end,
                        "参数来源": "; ".join(f"{k}: {v}" for k, v in provenance.items())})
            summary.append(row)
    preferred = []
    for metric in ("TTFT", "TPOT", "ITL", "E2EL"):
        preferred += [c for c in all_columns if c.startswith(f"请求/{metric}/") and ("/Average " in c or "/P99 " in c)]
    columns = SUMMARY_BASE + preferred + [c for c in all_columns if c not in preferred]
    columns += ["任务", "来源文件", "文件内结果序号", "起始行", "结束行", "参数来源"]
    return columns, summary, details


def write_excel(output, columns, rows, details, issues, force=False):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.table import Table, TableStyleInfo
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise ValueError("缺少 openpyxl，请执行: python -m pip install 'openpyxl>=3.1,<4'") from exc
    if output.exists() and not force:
        raise ValueError(f"输出文件已存在: {output}。使用 --force 覆盖")
    if max(len(rows), len(details), len(issues)) > 1048575 or len(columns) > 16384:
        raise ValueError("数据超过 Excel 单表容量，请按日志目录分批导出")
    workbook = Workbook()
    workbook.remove(workbook.active)

    def sheet(title, headers, values, freeze):
        ws = workbook.create_sheet(title)
        ws.append(headers)
        for record in values:
            ws.append(record)
        ws.freeze_panes = freeze
        ws.sheet_view.showGridLines = False
        ws.row_dimensions[1].height = 44
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="18364D")
            cell.font = Font(name="Calibri", size=11, color="FFFFFF", bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for cells in ws.iter_rows(min_row=2):
            for cell in cells:
                # Logs and user metadata are literal text, never Excel formulas.
                if isinstance(cell.value, str):
                    cell.value = re.sub(r"[\x00-\x08\x0b-\x0c\x0e-\x1f]", "", cell.value)[:32767]
                    cell.data_type = "s"
                cell.font = Font(name="Calibri", size=11, color="223746")
                cell.alignment = Alignment(vertical="top")
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '#,##0.####'
                if headers[cell.column - 1].endswith(" (%)"):
                    cell.number_format = "0.00%"
        for index, header_name in enumerate(headers, 1):
            if header_name in {"来源文件", "参数来源", "提示", "原始文本", "任务"}:
                width = 52
            elif "/" in header_name:
                width = 24
            else:
                width = max(12, min(30, len(header_name) * 1.8 + 3))
            ws.column_dimensions[get_column_letter(index)].width = width
        if ws.max_row > 1:
            table = Table(displayName=f"Report{len(workbook.worksheets)}", ref=ws.dimensions)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(table)
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_title_rows = "1:1"
        return ws

    sheet("汇总", columns, ([r.get(c) for c in columns] for r in rows), "F2")
    sheet("指标明细", ["结果ID", "阶段", "表格类型", "指标", "统计项", "数值", "单位", "原始文本", "来源文件", "行号", "提示"], details, "F2")
    if issues:
        sheet("解析提示", ["来源文件", "行号", "提示"], issues, "A2")
    output.parent.mkdir(parents=True, exist_ok=True)
    # Exclusive create prevents an accidental overwrite even if another process creates it mid-run.
    with output.open("wb" if force else "xb") as stream:
        workbook.save(stream)


def collect_inputs(inputs, recursive):
    paths, seen = [], set()
    for name in inputs:
        if name == "-":
            if name not in seen:
                seen.add(name)
                paths.append(name)
            continue
        matches = [Path(p) for p in glob.glob(name, recursive=recursive)]
        if not matches:
            raise ValueError(f"没有匹配的输入文件: {name}")
        for match in sorted(matches):
            if match.is_dir():
                files = sorted(match.rglob("*") if recursive else match.glob("*"))
                files = [p for p in files if p.is_file() and p.suffix.lower() in {".log", ".txt", ".out"}]
            else:
                files = [match]
            for path in files:
                key = str(path.resolve())
                if key not in seen:
                    seen.add(key)
                    paths.append(key)
    return paths


def read_text(path, encoding):
    data = Path(path).read_bytes()
    if encoding != "auto":
        return data.decode(encoding)
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("gb18030")


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("inputs", nargs="*", help="日志文件、通配符、目录或 - (stdin)")
    parser.add_argument("-o", "--output", type=Path, required=True, help="输出 .xlsx 路径")
    parser.add_argument("-r", "--recursive", action="store_true", help="递归扫描目录")
    parser.add_argument("--encoding", default="auto", help="auto: UTF-8 / BOM UTF-16 / GB18030")
    for flag in ("concurrency", "input-len", "output-len", "model", "dataset", "tag"):
        parser.add_argument("--" + flag, help="所有输入的默认元数据；每块 AISBENCH_META 可覆盖")
    parser.add_argument("--strict", action="store_true", help="存在解析提示则报错，不写 Excel")
    parser.add_argument("--force", action="store_true", help="覆盖已有输出文件")
    args = parser.parse_args(argv)
    try:
        if args.output.suffix.lower() != ".xlsx":
            raise ValueError("输出路径须以 .xlsx 结尾")
        inputs = args.inputs or (["-"] if not sys.stdin.isatty() else [])
        if not inputs:
            raise ValueError("请提供日志文件，或通过管道输入终端输出")
        paths = collect_inputs(inputs, args.recursive)
        if not paths:
            raise ValueError("目录中没有 .log/.txt/.out 文件")
        if str(args.output.resolve()) in paths:
            raise ValueError("输出路径不能与输入文件相同")
        defaults = {key: getattr(args, key) for key in META_FIELDS if getattr(args, key) is not None}
        runs, issues = [], []
        for path in paths:
            source = "<stdin>" if path == "-" else path
            text = sys.stdin.read() if path == "-" else read_text(path, args.encoding)
            parsed, warnings = parse_text(text, source, defaults)
            runs.extend(parsed)
            issues.extend(warnings)
        if not runs:
            raise ValueError("没有识别到任何结果。请保存包含 Performance Parameters / Common Metric 的完整表格")
        columns, rows, details = build_rows(runs, issues)
        if args.strict and issues:
            for source, line, message in issues[:12]:
                print(f"{source}:{line}: {message}", file=sys.stderr)
            raise ValueError(f"--strict: 共 {len(issues)} 条解析提示，未写入 Excel")
        write_excel(args.output, columns, rows, details, issues, args.force)
        print(f"已导出 {len(runs)} 个结果块、{len(rows)} 行结果/阶段、{len(details)} 个指标值 -> {args.output.resolve()}")
        if issues:
            print(f"{len(issues)} 条提示见 Excel 的“解析提示”工作表", file=sys.stderr)
        return 0
    except (OSError, UnicodeError, ValueError, LookupError) as exc:
        print(f"错误: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
